import tkinter as tk
from tkinter import messagebox, ttk
from pathlib import Path
from datetime import datetime
import os
import json
import time

# Import para sa Excel Automation
try:
    from openpyxl import load_workbook, Workbook
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False

# Import para sa QR Code Scanning
try:
    import cv2
    HAS_QR_DEPENDENCIES = True
except ImportError:
    cv2 = None
    HAS_QR_DEPENDENCIES = False

# Import para sa modernong Login Assets (Mula sa PIL)
from PIL import Image, ImageTk

# --------------------------
# CONFIGURATIONS & DATA MANAGEMENT
# --------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ATTENDANCE_FILE = os.path.join(BASE_DIR, "coepro4_attendance_data.json")
EXCEL_FILE = os.path.join(BASE_DIR, "Attendance_Records.xlsx")

ADMIN_USERNAME = "Admin"
ADMIN_PASSWORD = "coepro4"

# Login Window UI Constants
BASE_W = 1024
BASE_H = 768
BG = "#010912"
BG_DARK = "#00050c"
PANEL = "#061f2d"
INPUT = "#0d2a3a"
INPUT_FOCUS = "#123f55"
TEXT = "#eef8fb"
TEXT_SOFT = "#88aab4"
ACCENT = "#00d99a"
CYAN = "#16c9ff"
ASSETS = Path(__file__).parent / "assets"

LOGO_PNG = ASSETS / "logo_transparent.png"
BACKGROUND_PNG = ASSETS / "background.png"


# --------------------------
# CORE ATTENDANCE SYSTEM LOGIC
# --------------------------
class AttendanceSystem:
    def __init__(self):
        self.students = []
        self.attendance = []
        self.load_data()
        self.init_excel()
        self.scan_cooldown = {}
        self.cooldown_duration = 10 

    def load_data(self):
        if os.path.exists(ATTENDANCE_FILE):
            try:
                with open(ATTENDANCE_FILE, "r") as f:
                    data = json.load(f)
                    self.students = data.get("students", [])
                    self.attendance = data.get("attendance", [])
            except Exception:
                self.students = []
                self.attendance = []

    def check_cooldown(self, student_id):
        current_time = time.time()      
        if student_id in self.scan_cooldown:
            elapsed_time = current_time - self.scan_cooldown[student_id]
            if elapsed_time < self.cooldown_duration:
                return True
        self.scan_cooldown[student_id] = current_time
        return False

    def save_data(self):
        data = {"students": self.students, "attendance": self.attendance}
        with open(ATTENDANCE_FILE, "w") as f:
            json.dump(data, f, indent=4)

    def init_excel(self):
        if not HAS_EXCEL:
            return
        if not os.path.exists(EXCEL_FILE):
            wb = Workbook()
            ws = wb.active
            ws.title = "Attendance Logs"
            ws.append(["Student ID", "Full Name", "Course", "Year Level", "Location", "Date & Time"])
            wb.save(EXCEL_FILE)

    def save_to_excel(self, student_id, name, course, level, location, timestamp):
        if not HAS_EXCEL:
            return False
        try:
            wb = load_workbook(EXCEL_FILE)
            ws = wb.active
            ws.append([student_id, name, course, level, location, timestamp])
            wb.save(EXCEL_FILE)
            return True
        except PermissionError:
            messagebox.showerror("Excel Error", "Hindi mai-save sa Excel! Paki-sarahan ang 'Attendance_Records.xlsx'.")
            return False
        except Exception:
            return False

    def add_student(self, student_id, name, course, level, location="Room 101"):
        if not student_id.strip() or not name.strip() or not course.strip() or not level.strip():
            return False, "All fields are required!"

        student_id = student_id.strip()
        name = name.strip()
        course = course.strip()
        level = level.strip()
        location = location.strip() if location.strip() else "Room 101"
        
        now = datetime.now()
        current_time = now.strftime("%Y-%m-%d %H:%M:%S")
        today_date = now.strftime("%Y-%m-%d")

        for student in self.students:
            existing_id = student.get("student_id", student.get("id", ""))
            reg_time = student.get("registered_at", "")
            if existing_id == student_id and reg_time.startswith(today_date):
                return False, f"ID {student_id} ay naka-present na ngayong araw!"

        new_student = {
            "student_id": student_id,
            "name": name,
            "course": course,
            "level": level,
            "location": location,
            "registered_at": current_time
        }
        
        self.students.append(new_student)
        self.save_data()
        self.save_to_excel(student_id, name, course, level, location, current_time)
        return True, f"Logged at {location}: {name}!"

    def get_room_counts(self):
        today_date = datetime.now().strftime("%Y-%m-%d")
        counts = {"Room 101": 0, "Room 102": 0, "Room 103": 0, "Room 104": 0}
        for student in self.students:
            reg_time = student.get("registered_at", "")
            loc = student.get("location", "Room 101")
            if reg_time.startswith(today_date) and loc in counts:
                counts[loc] += 1
        return counts

    def get_hourly_traffic(self):
        traffic = {"07:00": 0, "08:00": 0, "09:00": 0, "10:00": 0, "11:00": 0, "12:00": 0}
        for student in self.students:
            reg_time = student.get("registered_at", "")
            if reg_time:
                try:
                    time_obj = datetime.strptime(reg_time, "%Y-%m-%d %H:%M:%S")
                    hour_str = time_obj.strftime("%H:00")
                    if hour_str in traffic:
                        traffic[hour_str] += 1
                except ValueError:
                    continue
        return traffic

    def get_absenteeism_insights(self):
        student_records = {}
        for student in self.students:
            s_id = student.get("student_id", "Unknown")
            name = student.get("name", "Unknown")
            if s_id not in student_records:
                student_records[s_id] = {"name": name, "logs": 0}
            student_records[s_id]["logs"] += 1
        
        sorted_students = sorted(student_records.items(), key=lambda item: item[1]["logs"])
        return sorted_students[:3]


# --------------------------
# DASHBOARD WIDGETS
# --------------------------
class CampusMapDashboard(tk.LabelFrame):
    def __init__(self, parent, system):
        super().__init__(parent, text=" Live Campus Occupancy Heat Map ", fg="white", bg="#2b2b2b", font=("Arial", 11, "bold"))
        self.system = system
        
        self.canvas = tk.Canvas(self, width=540, height=190, bg="#1a1a1a", highlightthickness=0)
        self.canvas.pack(pady=5, padx=10)
        
        self.rooms_config = {
            "Room 101": {"coords": (20, 15, 130, 105), "text_pos": (75, 60)},
            "Room 102": {"coords": (150, 15, 260, 105), "text_pos": (205, 60)},
            "Room 103": {"coords": (280, 15, 390, 105), "text_pos": (335, 60)},
            "Room 104": {"coords": (410, 15, 520, 160), "text_pos": (465, 87)},
        }
        self.room_shapes = {}
        self.update_map()

    def update_map(self):
        self.canvas.delete("all")
        room_counts = self.system.get_room_counts()

        self.canvas.create_rectangle(20, 120, 390, 160, fill="#2d2d2d", outline="#444444", width=1)
        self.canvas.create_text(205, 140, text="✦ MAIN CORRIDOR ✦", fill="#777777", font=("Arial", 8, "bold", "italic"))

        for room_name, config in self.rooms_config.items():
            count = room_counts.get(room_name, 0)
            
            if count == 0:
                color, status, text_color = "#2e7d32", "Vacant", "white"
            elif 1 <= count <= 3:
                color, status, text_color = "#fbc02d", "Moderate", "black"
            else:
                color, status, text_color = "#c62828", "CROWDED", "white"

            x1, y1, x2, y2 = config["coords"]
            rect_id = self.canvas.create_rectangle(x1, y1, x2, y2, fill=color, outline="#ffffff", width=2, tags=room_name)
            self.room_shapes[room_name] = rect_id

            tx, ty = config["text_pos"]
            self.canvas.create_text(tx, ty - 15, text=room_name, fill=text_color, font=("Arial", 10, "bold"))
            self.canvas.create_text(tx, ty + 5, text=f"Pax: {count}", fill=text_color, font=("Arial", 9, "bold"))
            self.canvas.create_text(tx, ty + 22, text=status, fill=text_color, font=("Arial", 7, "bold", "underline"))

            self.canvas.tag_bind(room_name, "<Enter>", lambda e, r=room_name: self.canvas.itemconfig(self.room_shapes[r], outline="#0078D7", width=3))
            self.canvas.tag_bind(room_name, "<Leave>", lambda e, r=room_name: self.canvas.itemconfig(self.room_shapes[r], outline="#ffffff", width=2))


class AnalyticsDashboard(tk.LabelFrame):
    def __init__(self, parent, system):
        super().__init__(parent, text=" Actionable Insights Portal ", fg="#00ff00", bg="#2b2b2b", font=("Arial", 11, "bold"))
        self.system = system
        
        self.left_frame = tk.Frame(self, bg="#2b2b2b")
        self.left_frame.pack(side="left", fill="both", expand=True, padx=5, pady=5)
        
        self.right_frame = tk.Frame(self, bg="#2b2b2b")
        self.right_frame.pack(side="right", fill="both", expand=True, padx=5, pady=5)
        
        tk.Label(self.left_frame, text="⚠️ At-Risk Attendance Alerts", fg="white", bg="#2b2b2b", font=("Arial", 9, "bold")).pack(anchor="w", padx=5)
        self.absentee_box = tk.Frame(self.left_frame, bg="#1a1a1a", height=120, width=240)
        self.absentee_box.pack(fill="both", expand=True, padx=5, pady=5)
        self.absentee_box.pack_propagate(False)

        tk.Label(self.right_frame, text="📈 Gate Peak Arrival Traffic Tracker", fg="white", bg="#2b2b2b", font=("Arial", 9, "bold")).pack(anchor="w", padx=5)
        self.traffic_canvas = tk.Canvas(self.right_frame, width=260, height=120, bg="#1a1a1a", highlightthickness=0)
        self.traffic_canvas.pack(fill="both", expand=True, padx=5, pady=5)

        self.refresh_analytics()

    def refresh_analytics(self):
        for widget in self.absentee_box.winfo_children():
            widget.destroy()
            
        at_risk_students = self.system.get_absenteeism_insights()
        if not at_risk_students:
            lbl = tk.Label(self.absentee_box, text="All students are within\nhealthy attendance marks.", fg="#777777", bg="#1a1a1a", font=("Arial", 9, "italic"))
            lbl.place(relx=0.5, rely=0.5, anchor="center")
        else:
            for s_id, info in at_risk_students:
                row = tk.Frame(self.absentee_box, bg="#1a1a1a", pady=2)
                row.pack(fill="x", padx=5)
                tk.Label(row, text="●", fg="#dc3545", bg="#1a1a1a", font=("Arial", 10)).pack(side="left", padx=2)
                tk.Label(row, text=f"{info['name']} ({s_id})", fg="white", bg="#1a1a1a", font=("Arial", 9, "bold")).pack(side="left", padx=2)
                tk.Label(row, text=f"Logs: {info['logs']}d present", fg="#ffc107", bg="#1a1a1a", font=("Arial", 8, "italic")).pack(side="right", padx=5)

        self.traffic_canvas.delete("all")
        traffic_data = self.system.get_hourly_traffic()
        
        max_value = max(traffic_data.values()) if max(traffic_data.values()) > 0 else 1
        x_start = 25
        y_bottom = 100
        bar_width = 25
        spacing = 12

        for i, (hour, count) in enumerate(traffic_data.items()):
            bar_height = (count / max_value) * 70
            x1 = x_start + i * (bar_width + spacing)
            y1 = y_bottom - bar_height
            x2 = x1 + bar_width
            y2 = y_bottom
            
            color = "#dc3545" if count == max_value and count > 0 else "#0078D7"
            
            self.traffic_canvas.create_rectangle(x1, y1, x2, y2, fill=color, outline="")
            self.traffic_canvas.create_text((x1 + x2)/2, y1 - 8, text=str(count), fill="white", font=("Arial", 8, "bold"))
            self.traffic_canvas.create_text((x1 + x2)/2, y_bottom + 10, text=hour.split(":")[0], fill="#888888", font=("Arial", 8))


# --------------------------
# MAIN DASHBOARD WINDOW
# --------------------------
class CoeproAttendanceApp(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent
        self.title("COEPRO4 Smart Campus Dashboard")
        self.geometry("600x690")
        self.configure(bg="#1e1e1e")
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        
        self.system = AttendanceSystem()
        self.cap = None
        self.is_scanning = False

        title_lbl = tk.Label(self, text="Class Attendance & Spatial Dashboard", font=("Arial", 16, "bold"), fg="white", bg="#1e1e1e")
        title_lbl.pack(pady=10)

        self.map_dashboard = CampusMapDashboard(self, self.system)
        self.map_dashboard.pack(fill="x", padx=20, pady=2)

        self.analytics_portal = AnalyticsDashboard(self, self.system)
        self.analytics_portal.pack(fill="x", padx=20, pady=5)

        self.log_frame = tk.LabelFrame(self, text=" Live Tracking Terminal ", fg="white", bg="#2b2b2b", font=("Arial", 9, "bold"))
        self.log_frame.pack(fill="both", expand=True, padx=20, pady=5)
        
        self.log_text = tk.Text(self.log_frame, height=4, bg="#121212", fg="#00ff00", font=("Consolas", 9), state="disabled", highlightthickness=0)
        self.log_text.pack(fill="both", expand=True, padx=5, pady=5)

        btn_frame = tk.Frame(self, bg="#1e1e1e")
        btn_frame.pack(pady=10)

        self.qr_btn = tk.Button(btn_frame, text="📷 Start Real-time Scan", font=("Arial", 11, "bold"), bg="#0078D7", fg="white", padx=15, pady=6, command=self.toggle_scanner)
        self.qr_btn.grid(row=0, column=0, padx=10)

        refresh_btn = tk.Button(btn_frame, text="🔄 Refresh Systems", font=("Arial", 11, "bold"), bg="#6c757d", fg="white", padx=15, pady=6, command=self.global_refresh)
        refresh_btn.grid(row=0, column=1, padx=10)

    def append_log(self, message):
        self.log_text.configure(state="normal")
        timestamp = datetime.now().strftime("[%H:%M:%S] ")
        self.log_text.insert(tk.END, timestamp + message + "\n")
        self.log_text.see(tk.END)
        self.log_text.configure(state="disabled")

    def global_refresh(self):
        self.map_dashboard.update_map()
        self.analytics_portal.refresh_analytics()
        self.append_log("Global spatial mapping and analytic graphs refreshed.")

    def toggle_scanner(self):
        if self.is_scanning:
            self.stop_scanner()
        else:
            if not HAS_QR_DEPENDENCIES or cv2 is None:
                messagebox.showerror("Error", "Missing system architectural modules for OpenCV.")
                return
            self.cap = cv2.VideoCapture(0)
            if not self.cap.isOpened():
                messagebox.showerror("Camera Error", "Hardware pipeline error: Webcam cannot open.")
                return
            
            self.is_scanning = True
            self.qr_btn.configure(text="🛑 Stop Scanning", bg="#dc3545")
            self.append_log("Scanner listening online... Align QR payload matrix.")
            self.detector = cv2.QRCodeDetector()
            self.process_camera_frame()

    def stop_scanner(self):
        self.is_scanning = False
        self.qr_btn.configure(text="📷 Start Real-time Scan", bg="#0078D7")
        if self.cap:
            self.cap.release()
            self.cap = None
        cv2.destroyAllWindows()
        self.append_log("Scanner listening active mode closed.")

    def process_camera_frame(self):
        if not self.is_scanning:
            return

        ret, frame = self.cap.read()
        if ret:
            retval, decoded_info, _, _ = self.detector.detectAndDecodeMulti(frame)
            if retval and decoded_info:
                for qr_data in decoded_info:
                    if qr_data.strip():
                        self.handle_qr_payload(qr_data)
            
            cv2.putText(frame, "COEPRO4 Actionable Analytics Mode", (15, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 255, 255), 1)
            cv2.imshow("Smart Campus Camera View", frame)

        self.after(15, self.process_camera_frame)

    def handle_qr_payload(self, qr_data):
        s_id, s_name, s_course, s_level, s_loc = "", "", "", "", "Room 101"
        try:
            student_info = json.loads(qr_data)
            s_id = str(student_info.get("student_id", student_info.get("StudentID", "")))
            s_name = student_info.get("name", student_info.get("FullName", ""))
            s_course = student_info.get("course", student_info.get("Course", ""))
            s_level = student_info.get("level", student_info.get("YearLevel", ""))
            s_loc = student_info.get("location", "Room 101")
        except json.JSONDecodeError:
            tokens = [token.strip() for token in qr_data.split(",")]
            if len(tokens) >= 4:
                s_id, s_name, s_course, s_level = tokens[:4]
                if len(tokens) >= 5:
                    s_loc = tokens[4]

        if s_id and s_name:
            if self.system.check_cooldown(s_id):
                return
            success, msg = self.system.add_student(s_id, s_name, s_course, s_level, s_loc)
            if success:
                self.append_log(f"SUCCESS: {s_name} entered {s_loc}")
                self.map_dashboard.update_map()
                self.analytics_portal.refresh_analytics()
            else:
                self.append_log(f"REJECTED: {msg}")

    def on_close(self):
        self.stop_scanner()
        self.parent.destroy()


# --------------------------
# MODERN LOGIN WINDOW (MAIN ENGINE)
# --------------------------
class OneMinuteLogin(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("One Minute - Sign In")
        self.geometry(f"{BASE_W}x{BASE_H}")
        self.minsize(900, 650)
        self.configure(bg=BG)
        self.resizable(True, True)

        self.canvas = tk.Canvas(self, highlightthickness=0, bg=BG)
        self.canvas.pack(fill="both", expand=True)

        self.username_var = tk.StringVar()
        self.passcode_var = tk.StringVar()
        self.logo_photo = None
        self.logo_small = None
        self.background_photo = None
        self.entries = []
        self.resize_job = None
        self.login_button = None

        self.bind("<Configure>", self.schedule_redraw)
        self.bind("<F11>", self.toggle_fullscreen)
        self.bind("<Escape>", self.exit_fullscreen)
        self.bind("<Return>", lambda event: self.login())
        self.after(80, self.maximize_window)
        self.redraw()

    def maximize_window(self):
        try:
            self.state("zoomed")
        except tk.TclError:
            self.attributes("-fullscreen", True)

    def toggle_fullscreen(self, event=None):
        self.attributes("-fullscreen", not self.attributes("-fullscreen"))

    def exit_fullscreen(self, event=None):
        self.attributes("-fullscreen", False)

    def schedule_redraw(self, event=None):
        if event is not None and event.widget is not self:
            return
        if self.resize_job:
            self.after_cancel(self.resize_job)
        self.resize_job = self.after(70, self.redraw)

    def redraw(self):
        self.resize_job = None
        for entry in self.entries:
            entry.destroy()
        self.entries.clear()
        self.canvas.delete("all")
        self.canvas.config(width=self.winfo_width(), height=self.winfo_height())

        self.draw_background()
        self.draw_top_logo()
        self.draw_signin_form()
        self.draw_waves()
        self.draw_footer()

    def layout(self):
        w = max(self.winfo_width(), BASE_W)
        h = max(self.winfo_height(), BASE_H)
        scale = min(w / BASE_W, h / BASE_H)
        offset_x = (w - BASE_W * scale) / 2
        return w, h, scale, offset_x

    def x(self, value):
        _, _, scale, offset_x = self.layout()
        return offset_x + value * scale

    def y(self, value):
        _, _, scale, _ = self.layout()
        return value * scale

    def s(self, value):
        _, _, scale, _ = self.layout()
        return value * scale

    def font(self, family, size, weight=None):
        size = max(8, int(self.s(size)))
        return (family, size, weight) if weight else (family, size)

    def ui_font(self, size, weight=None):
        return self.font("Century Gothic", size, weight)

    def draw_background(self):
        w, h, _, _ = self.layout()
        self.canvas.create_rectangle(0, 0, w, h, fill=BG, outline="")
        if BACKGROUND_PNG.exists():
            img = Image.open(BACKGROUND_PNG)
            scale = max(w / img.width, h / img.height)
            new_w = int(img.width * scale)
            new_h = int(img.height * scale)
            img = img.resize((new_w, new_h), Image.LANCZOS).convert("RGB")
            dark_layer = Image.new("RGB", img.size, BG)
            img = Image.blend(img, dark_layer, 0.72)
            self.background_photo = ImageTk.PhotoImage(img)
            x = (w - new_w) // 2
            y = (h - new_h) // 2
            self.canvas.create_image(x, y, image=self.background_photo, anchor="nw")

    def draw_top_logo(self):
        if LOGO_PNG.exists():
            logo = Image.open(LOGO_PNG)
            logo_w = max(220, int(self.s(250)))
            logo_h = max(1, int(logo.height * (logo_w / logo.width)))
            logo = logo.resize((logo_w, logo_h), Image.LANCZOS)
            self.logo_photo = ImageTk.PhotoImage(logo)
            self.logo_small = self.logo_photo
            self.canvas.create_image(self.x(512), self.y(18), image=self.logo_small, anchor="n")
            return

        self.draw_mark(440, 52, 0.42)
        self.canvas.create_text(self.x(470), self.y(53), text="ONE MINUTE", fill=TEXT,
                                font=self.ui_font(17, "bold"), anchor="w")

    def draw_mark(self, cx, cy, factor=1.0):
        r = self.s(18 * factor)
        cx = self.x(cx)
        cy = self.y(cy)
        self.canvas.create_arc(cx - r, cy - r, cx + r, cy + r, start=130, extent=250,
                               outline=TEXT, width=max(2, int(self.s(5 * factor))))
        self.canvas.create_arc(cx - r, cy - r, cx + r, cy + r, start=280, extent=70,
                               outline=CYAN, width=max(2, int(self.s(5 * factor))))
        self.canvas.create_line(cx, cy, cx + self.s(12 * factor), cy - self.s(16 * factor),
                                fill=ACCENT, width=max(2, int(self.s(4 * factor))))
        self.canvas.create_oval(cx - self.s(4 * factor), cy - self.s(4 * factor),
                                cx + self.s(4 * factor), cy + self.s(4 * factor),
                                fill=ACCENT, outline="")

    def draw_signin_form(self):
        center_x = 512
        self.canvas.create_text(self.x(center_x), self.y(170), text="Sign in", fill=TEXT,
                                font=self.ui_font(32), anchor="center")
        self.canvas.create_text(self.x(center_x), self.y(221),
                                text="Welcome back! Sign in to continue",
                                fill=TEXT_SOFT, font=self.ui_font(12), anchor="center")

        self.add_input("username", self.username_var, 364, 285, False)
        self.add_input("passcode", self.passcode_var, 364, 342, True)
        self.add_login_button()

    def add_input(self, placeholder, variable, x, y, secret):
        width = 296
        height = 38
        self.rounded_rectangle(self.x(x), self.y(y), self.x(x + width), self.y(y + height),
                               radius=self.s(8), fill=INPUT, outline="#2a6673")
        entry = tk.Entry(
            self,
            textvariable=variable,
            show="",
            bg=INPUT,
            fg="#a8c4c9",
            insertbackground=TEXT,
            relief="flat",
            borderwidth=0,
            font=self.ui_font(11),
        )
        if variable.get() == "":
            variable.set(placeholder)
        elif secret and variable.get() != placeholder:
            entry.config(show="*")

        def on_focus_in(event):
            if entry.get() == placeholder:
                entry.delete(0, "end")
                entry.config(fg=TEXT)
                if secret:
                    entry.config(show="*")
            self.canvas.itemconfig(entry.bg_shape, fill=INPUT_FOCUS, outline="#328395")
            entry.config(bg=INPUT_FOCUS)

        def on_focus_out(event):
            if not entry.get():
                variable.set(placeholder)
                entry.config(show="")
                entry.config(fg="#a8c4c9")
            self.canvas.itemconfig(entry.bg_shape, fill=INPUT, outline="#2a6673")
            entry.config(bg=INPUT)

        entry.bg_shape = self.canvas.find_all()[-1]
        entry.bind("<FocusIn>", on_focus_in)
        entry.bind("<FocusOut>", on_focus_out)
        self.entries.append(entry)
        self.canvas.create_window(self.x(x + 14), self.y(y + 19), anchor="w",
                                  width=self.s(width - 28), height=self.s(24), window=entry)

    def add_login_button(self):
        x = 364
        y = 420
        width = 296
        height = 42
        self.login_button = self.rounded_rectangle(self.x(x), self.y(y), self.x(x + width), self.y(y + height),
                                                   radius=self.s(10), fill=ACCENT, outline=ACCENT)
        self.login_text = self.canvas.create_text(self.x(x + width / 2), self.y(y + height / 2),
                                                  text="Login", fill="#06433a",
                                                  font=self.ui_font(12, "bold"))
        for item in (self.login_button, self.login_text):
            self.canvas.tag_bind(item, "<Enter>", self.on_button_enter)
            self.canvas.tag_bind(item, "<Leave>", self.on_button_leave)
            self.canvas.tag_bind(item, "<Button-1>", lambda event: self.login())

    def draw_waves(self):
        w, h, _, _ = self.layout()
        y1 = self.y(560)
        y2 = self.y(600)
        y3 = self.y(640)
        self.canvas.create_polygon(
            0, y1, self.x(160), self.y(575), self.x(350), self.y(570), self.x(565), self.y(540),
            self.x(770), self.y(532), w, self.y(548), w, h, 0, h,
            fill="#041f31", outline=""
        )
        self.canvas.create_polygon(
            0, y2, self.x(190), self.y(622), self.x(390), self.y(612), self.x(620), self.y(584),
            self.x(810), self.y(568), w, self.y(585), w, h, 0, h,
            fill="#031725", outline=""
        )
        self.canvas.create_polygon(
            0, y3, self.x(180), self.y(682), self.x(390), self.y(672), self.x(625), self.y(642),
            self.x(820), self.y(632), w, self.y(654), w, h, 0, h,
            fill="#010a14", outline=""
        )

    def draw_footer(self):
        self.canvas.create_text(self.x(512), self.y(724), text="2026 | One Minute. All rights reserved.",
                                fill="#5d7c84", font=self.ui_font(8), anchor="center")

    def login(self):
        username = self.username_var.get().strip()
        passcode = self.passcode_var.get().strip()

        if username in ("", "username") or passcode in ("", "passcode"):
            messagebox.showwarning("Missing Details", "Please enter your username and passcode.")
        elif username == ADMIN_USERNAME and passcode == ADMIN_PASSWORD:
            messagebox.showinfo("Login Successful", "Welcome, Admin!")
            self.withdraw()  # Itatago ang Login Window
            
            # Buksan ang pangunahing COEPRO Dashboard Window
            dashboard = CoeproAttendanceApp(self)
            dashboard.grab_set()
        else:
            messagebox.showerror("Login Failed", "Invalid username or passcode.")

    def on_button_enter(self, event=None):
        self.canvas.itemconfig(self.login_button, fill="#31f29e", outline="#31f29e")

    def on_button_leave(self, event=None):
        self.canvas.itemconfig(self.login_button, fill=ACCENT, outline=ACCENT)

    def rounded_rectangle(self, x1, y1, x2, y2, radius=20, **kwargs):
        radius = int(radius)
        points = [
            x1 + radius, y1,
            x2 - radius, y1,
            x2, y1,
            x2, y1 + radius,
            x2, y2 - radius,
            x2, y2,
            x2 - radius, y2,
            x1 + radius, y2,
            x1, y2,
            x1, y2 - radius,
            x1, y1 + radius,
            x1, y1,
        ]
        return self.canvas.create_polygon(points, smooth=True, **kwargs)


if __name__ == "__main__":
    app = OneMinuteLogin()
    app.mainloop()
